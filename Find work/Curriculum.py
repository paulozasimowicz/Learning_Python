# Run:
# pip install python-docx
# python make_cv.py

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/paulovitor/Downloads/CV_STFI_Calaca.docx"

ACCENT = "1F4E79"
LINE = "9EADCC"
TEXT = "222222"
MUTED = "555555"


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = f"w:{edge}"
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)

            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_table_no_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                left={"val": "nil"},
                right={"val": "nil"},
                insideH={"val": "nil"},
                insideV={"val": "nil"},
            )


def add_hyperlink(paragraph, text, url, color="0563C1", underline=True):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rPr.append(c)

    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_run(run, bold=False, italic=False, size=None, color=None):
    run.bold = bold
    run.italic = italic

    if size is not None:
        run.font.size = Pt(size)

    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)

    return run


def compact(p, before=0, after=0, line_spacing=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    return p


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    compact(p, before=6, after=1)

    r = p.add_run(title.upper())
    set_run(r, bold=True, size=9.2, color=ACCENT)

    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LINE)

    pBdr.append(bottom)
    pPr.append(pBdr)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    compact(p, before=0, after=0, line_spacing=1.0)
    p.paragraph_format.left_indent = Inches(0.18 + level * 0.15)
    p.paragraph_format.first_line_indent = Inches(-0.12)

    r = p.add_run(text)
    set_run(r, size=8.15, color=TEXT)


def add_entry(doc, role, org="", location="", dates="", details=None, bullets=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_no_borders(table)
    table.autofit = False

    table.columns[0].width = Inches(4.75)
    table.columns[1].width = Inches(2.05)

    left, right = table.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    p = left.paragraphs[0]
    compact(p, after=0)

    set_run(p.add_run(role), bold=True, size=8.8, color="000000")

    if org:
        set_run(p.add_run("\n" + org), bold=True, size=8.0, color="444444")

    if details:
        set_run(p.add_run("\n" + details), italic=True, size=7.8, color=MUTED)

    p2 = right.paragraphs[0]
    compact(p2, after=0)
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    if location:
        set_run(p2.add_run(location), bold=True, size=8.0, color="000000")

    if dates:
        set_run(p2.add_run(("\n" if location else "") + dates), size=8.0, color=TEXT)

    for b in bullets or []:
        add_bullet(doc, b)


def add_kv(doc, key, value):
    p = doc.add_paragraph()
    compact(p, before=0, after=0, line_spacing=1.0)

    set_run(p.add_run(key + ": "), bold=True, size=8.15, color="000000")
    set_run(p.add_run(value), size=8.15, color=TEXT)


def add_small_para(doc, text):
    p = doc.add_paragraph()
    compact(p, before=0, after=1, line_spacing=1.02)
    set_run(p.add_run(text), size=8.3, color=TEXT)


doc = Document()

sec = doc.sections[0]
sec.top_margin = Inches(0.48)
sec.bottom_margin = Inches(0.48)
sec.left_margin = Inches(0.55)
sec.right_margin = Inches(0.55)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(8.4)
styles["Normal"].paragraph_format.space_after = Pt(0)

try:
    styles["List Bullet"].font.name = "Arial"
    styles["List Bullet"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["List Bullet"].font.size = Pt(8.1)
except Exception:
    pass


# Header
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(p, after=0)
set_run(
    p.add_run("Paulo Vitor Zasimowicz Pinto Calaça"),
    bold=True,
    size=16.2,
    color="000000",
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(p, after=2)

set_run(p.add_run("pvzpcp@gmail.com | +32 467 87 86 24 | "), size=8.1, color=TEXT)
add_hyperlink(p, "LinkedIn", "https://www.linkedin.com/in/paulocalaca")
set_run(p.add_run(" | "), size=8.1, color=TEXT)
add_hyperlink(p, "GitHub", "https://github.com/paulozasimowicz")


add_section_heading(doc, "Profile")
add_small_para(
    doc,
    "Textile engineering MSc candidate with thesis defended and formal graduation expected in September 2026, "
    "and a chemical engineering background. Research experience in 3D/4D body-scan data, computational geometry, "
    "mesh analysis and Python-based scientific tool development for textile applications. Combines textile-engineering "
    "coursework with process-engineering fundamentals, empirical data evaluation and industrial experience in data-driven "
    "process monitoring and coordination. Interested in technical textiles, textile lightweight construction, sustainable "
    "material systems and applied R&D with industry partners.",
)


add_section_heading(doc, "Education")
add_entry(
    doc,
    role="International MSc in Textile Engineering",
    org="Universiteit Ghent - University of West Attica - University of Borås - Technische Universität Dresden",
    dates="Sep 2024 - Sep 2026",
    details=(
        "Thesis: Accuracy Evaluation of 3D and 4D Body Scans | Thesis defended | "
        "Supervisors: Prof. Yordan Kyosev and Anselm Naake"
    ),
    bullets=[
        "Relevant coursework: Advanced Fibre Technology; Advanced and Specialised Textile Processing - Mechanical; "
        "Mechanics of Textile Materials; Polymer Technology; Biomaterials; Nanotechnology in the Textile Branch; "
        "Sustainable Textile Design; Garment Technology; Virtual Product Development.",
        "Academic focus on technical textiles, textile processes, garment and product development, computational methods and sustainable textile systems.",
    ],
)

add_entry(
    doc,
    role="Bachelor's in Chemical Engineering",
    org="Universidade Federal de Uberlândia, Brazil",
    location="Uberlândia, Brazil",
    dates="Graduated Feb 2024",
    bullets=[
        "Engineering background in process modelling and simulation, transport phenomena, chemical thermodynamics, "
        "unit operations, materials/mechanics, safety and industrial process analysis."
    ],
)


add_section_heading(doc, "Research and technical experience")
add_entry(
    doc,
    role="Master Thesis Research - 3D/4D Body Scan Accuracy for Textile Applications",
    org="Technische Universität Dresden, Institute of Textile Machinery and High Performance Material Technology",
    location="Dresden, Germany",
    dates="2025 - 2026",
    bullets=[
        "Developed a Python-based workflow for geometric accuracy evaluation of static and dynamic 3D/4D body-scan data relevant to textile product development.",
        "Processed surface meshes and scan sequences to evaluate girth measurements, temporal deformation, motion-induced geometric variation and deviations from manual reference measurements.",
        "Built interactive visualization and analysis tools using Python, vedo, Open3D, PyQt, NumPy, Matplotlib and scikit-learn.",
        "Prepared empirical data and visual outputs for scientific reporting, including method limitations and textile-relevant interpretation.",
    ],
)

add_entry(
    doc,
    role="Research Intern",
    org="Technische Universität Dresden",
    location="Dresden, Germany",
    dates="Jul 2025 - Aug 2025",
    bullets=[
        "Conducted research at the Institute of Textile Machinery and High Performance Material Technology.",
        "Built interactive visualization tools in Python to streamline point selection, filtering and analysis of complex scan data.",
        "Processed and analysed 4D scan data using visualization, clustering and computational geometry techniques.",
    ],
)


add_section_heading(doc, "Professional experience")
add_entry(
    doc,
    role="Human Resources Management Intern",
    org="Procter & Gamble",
    location="Rio de Janeiro, Brazil",
    dates="Jun 2023 - Nov 2023",
    bullets=[
        "Coordinated internal communication activities and organized more than 40 leadership meetings with plant leadership and directors.",
        "Co-created workshops and events with senior stakeholders, strengthening project coordination, communication and organizational skills.",
    ],
)

add_entry(
    doc,
    role="Environmental Intern",
    org="Coca-Cola S/A",
    location="Uberlândia, Brazil",
    dates="Feb 2022 - May 2022",
    bullets=[
        "Developed daily dashboards in Excel and Power BI for water-management indicators and operational monitoring.",
        "Detected and corrected three major water-consumption inefficiencies, supporting resource optimization and process improvement.",
    ],
)


doc.add_page_break()


add_section_heading(doc, "Additional experience")
add_entry(
    doc,
    role="Disney International Program",
    org="The Walt Disney Company",
    location="Orlando, United States",
    dates="Dec 2022 - Mar 2023",
    bullets=[
        "Worked in fast-paced Quick Service Food & Beverage operations and provided guest service in a multicultural environment."
    ],
)


add_section_heading(doc, "Leadership and extracurricular activities")
add_entry(
    doc,
    role="Member of the Collegiate - Faculty of Chemical Engineering",
    org="Universidade Federal de Uberlândia",
    location="Uberlândia, Brazil",
    dates="Feb 2021 - Mar 2023",
    bullets=[
        "Served as student representative for a course with more than 600 enrolled students.",
        "Collaborated with the Dean and faculty leadership on registration, deadline extensions, internships and academic matters.",
        "Supported didactic orientation, supervision and coordination of student-related topics.",
    ],
)


add_section_heading(doc, "Technical skills")
add_kv(doc, "Programming and data analysis", "Python, NumPy, Matplotlib, scikit-learn, Excel, Power BI, PowerPoint.")
add_kv(
    doc,
    "3D/4D analysis",
    "Mesh processing, surface geometry analysis, 3D and 4D human-shape analysis, computational body modelling, "
    "interactive scientific visualization, scan-data filtering and clustering.",
)
add_kv(
    doc,
    "Textile and materials knowledge",
    "Technical textiles, advanced fibre technology, high-performance textile materials, textile processing, garment technology, "
    "sustainable textile design, biomaterials, polymer technology and nanotechnology in textiles.",
)
add_kv(
    doc,
    "Engineering background",
    "Chemical engineering, process modelling and simulation, unit operations, transport phenomena, resource monitoring and process optimization.",
)


add_section_heading(doc, "Languages")
add_kv(doc, "Portuguese", "Native")
add_kv(doc, "English", "Fluent")
add_kv(doc, "Spanish", "Intermediate")
add_kv(doc, "French", "Basic")
add_kv(doc, "German", "Basic")


add_section_heading(doc, "Certificates")
add_kv(doc, "Green Belt Six Sigma Training", "Process improvement and structured problem solving.")
add_kv(doc, "ELA254: AIChE's Engineering Leadership Development", "Leadership development for engineering contexts.")


add_section_heading(doc, "Targeted fit for STFI lightweight construction role")
add_small_para(
    doc,
    "Relevant strengths for the STFI role include applied research experience, empirical data evaluation, Python-based method development, "
    "textile-engineering coursework, chemical/process-engineering fundamentals, and experience coordinating with academic and industrial-style stakeholders. "
    "Academic preparation in advanced fibre technology, biomaterials, polymer technology and sustainable textile design supports a strong interest in natural-fibre "
    "and biopolymer-based lightweight material systems.",
)


for s in doc.sections:
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    compact(footer, after=0)

    r = footer.add_run("Paulo Vitor Zasimowicz Pinto Calaça - Curriculum Vitae")
    r.font.size = Pt(7)
    r.font.color.rgb = RGBColor(120, 120, 120)


core = doc.core_properties
core.author = "Paulo Vitor Zasimowicz Pinto Calaça"
core.title = "Curriculum Vitae - STFI"
core.subject = "CV for STFI application"
core.comments = "Generated from recent CV and updated after thesis defense."

doc.save(OUT)
print(OUT)