# Run:
# pip install python-docx
# python make_letter.py
#
# Output:
# Motivation_STFI_Calaca.docx

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ---------------------------------------------------------------------
# Output path
# ---------------------------------------------------------------------

OUT = Path.home() / "Downloads" / "Motivation_STFI_Calaca.docx"


# ---------------------------------------------------------------------
# Editable information
# ---------------------------------------------------------------------

FULL_NAME = "Paulo Vitor Zasimowicz Pinto Calaça"
EMAIL = "pvzpcp@gmail.com"
PHONE = "+32 467 87 86 24"
LINKEDIN_URL = "https://www.linkedin.com/in/paulocalaca"
GITHUB_URL = "https://github.com/paulozasimowicz"

PLACE_AND_DATE = "Dresden, 23 June 2026"

RECIPIENT_LINES = [
    "Saxon Textile Research Institute (STFI)",
    "Recruitment Team",
    "Chemnitz, Germany",
]

SUBJECT = "Application for Research Associate in Textile Lightweight Construction"

BODY_PARAGRAPHS = [
    "Dear STFI Recruitment Team,",

    (
        "I am writing to apply for the position of Research Associate in the field of textile lightweight construction. "
        "I am currently completing the International MSc in Textile Engineering, with my master’s thesis already defended "
        "and formal graduation expected in September 2026. Before this, I completed a Bachelor’s degree in Chemical "
        "Engineering, which gave me a strong foundation in engineering thinking, materials, processes and data evaluation."
    ),

    (
        "What attracted me to this position is that STFI works close to real textile applications and industrial problems. "
        "During my thesis at Technische Universität Dresden, supervised by Prof. Yordan Kyosev and Dipl.-Wi.-Ing. Anselm "
        "Naake, I became especially interested in how research methods can be turned into practical tools. My work focused "
        "on 3D and 4D body-scan analysis for textile applications, where I developed Python-based workflows for visualizing, "
        "cleaning and evaluating scan data. This made me interested in research work where software, experiments and textile "
        "engineering are connected."
    ),

    (
        "I know that my background is not yet specifically in natural-fibre reinforced plastics or aviation applications. "
        "However, I believe I could contribute through my ability to learn technical topics, structure data, build analysis "
        "tools and create software workflows that support research work. For example, I would be interested in helping "
        "develop programs that can serve as a first computational baseline for evaluating or predicting lightweight textile "
        "structures. I enjoy working on problems where the method is not fully ready yet and needs to be built, tested and "
        "improved step by step."
    ),

    (
        "In addition to my research experience, I have worked in industrial environments at Procter & Gamble and Coca-Cola, "
        "where I developed communication, coordination and problem-solving skills. I am also aware that my German level does "
        "not yet meet the requested C1 level. I am currently taking German lessons at the Goethe-Institut and am at level "
        "A2.2, with the intention to continue improving. I would be happy to bring my textile engineering background, "
        "programming experience and motivation to learn into STFI’s research environment."
    ),

    "Thank you for considering my application. I would be pleased to discuss my background and motivation in an interview.",

    "Sincerely,\nPaulo Vitor Zasimowicz Pinto Calaça",
]

# ---------------------------------------------------------------------
# Style settings: same visual logic as the CV
# ---------------------------------------------------------------------

FONT_NAME = "Arial"
TEXT = "222222"
MUTED = "555555"
ACCENT = "1F4E79"


def set_font_for_style(style, font_name=FONT_NAME, size=10):
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style.font.size = Pt(size)


def set_run(run, bold=False, italic=False, size=None, color=None):
    run.bold = bold
    run.italic = italic

    if size is not None:
        run.font.size = Pt(size)

    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)

    return run


def compact(paragraph, before=0, after=0, line_spacing=1.0):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    return paragraph


def add_hyperlink(paragraph, text, url, color="0563C1", underline=True):
    """
    Adds a clickable hyperlink to a python-docx paragraph.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rPr.append(c)

    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


# ---------------------------------------------------------------------
# Document creation
# ---------------------------------------------------------------------

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
set_font_for_style(styles["Normal"], FONT_NAME, 10)


# Header: name
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(p, after=1)

set_run(
    p.add_run(FULL_NAME),
    bold=True,
    size=15.5,
    color="000000",
)


# Contact line
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(p, after=8)

set_run(p.add_run(f"{EMAIL} | {PHONE} | "), size=9, color=TEXT)
add_hyperlink(p, "LinkedIn", LINKEDIN_URL)
set_run(p.add_run(" | "), size=9, color=TEXT)
add_hyperlink(p, "GitHub", GITHUB_URL)


# Place and date
p = doc.add_paragraph()
compact(p, after=12)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(p.add_run(PLACE_AND_DATE), size=10, color=TEXT)


# Recipient
for line in RECIPIENT_LINES:
    p = doc.add_paragraph()
    compact(p, after=0)
    set_run(p.add_run(line), size=10, color=TEXT)


# Subject
p = doc.add_paragraph()
compact(p, before=12, after=8)

set_run(p.add_run("Subject: "), bold=True, size=10, color="000000")
set_run(p.add_run(SUBJECT), bold=True, size=10, color=ACCENT)


# Body
for i, paragraph_text in enumerate(BODY_PARAGRAPHS):
    p = doc.add_paragraph()

    if i == 0:
        compact(p, before=0, after=8, line_spacing=1.08)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif i == len(BODY_PARAGRAPHS) - 1:
        compact(p, before=6, after=0, line_spacing=1.08)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        compact(p, before=0, after=8, line_spacing=1.08)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    set_run(p.add_run(paragraph_text), size=10, color=TEXT)


# Footer
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
compact(footer, after=0)

r = footer.add_run("Paulo Vitor Zasimowicz Pinto Calaça - Motivation Letter")
r.font.size = Pt(7)
r.font.color.rgb = RGBColor(120, 120, 120)


# Metadata
core = doc.core_properties
core.author = FULL_NAME
core.title = "Motivation Letter - STFI"
core.subject = "Application for Research Associate in Textile Lightweight Construction"
core.comments = ""


# Save document
doc.save(OUT)
print(f"Saved: {OUT}")